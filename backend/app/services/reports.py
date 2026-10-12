import csv
import io
from datetime import UTC, datetime

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models import Event, Membership, Order, OrderItem, Pass, PassScan, PassTier, User
from app.models.enums import OrderStatus, PassStatus
from app.schemas.reports import DayRow, EventReport, ReportTotals, TierRow, WingRow
from app.services.events import tier_usage
from app.services.scanning import day_stats

APPROVED, PENDING = OrderStatus.APPROVED.value, OrderStatus.PENDING.value


async def build_report(session: AsyncSession, ev: Event) -> EventReport:
    # --- totals (one grouped query over orders)
    rows = (await session.execute(
        select(Order.status, Order.payment_submitted_at.is_not(None), func.count(), func.coalesce(func.sum(Order.total_amount), 0))
        .where(Order.event_id == ev.id).group_by(Order.status, Order.payment_submitted_at.is_not(None))
    )).all()
    t = dict(collected_amount=0, pending_amount=0, pending_count=0, awaiting_payment_count=0, approved_orders=0, rejected_orders=0)
    for status, submitted, n, amount in rows:
        if status == APPROVED:
            t["approved_orders"] += n
            t["collected_amount"] += int(amount)
        elif status == PENDING:
            if submitted:
                t["pending_count"] += n
                t["pending_amount"] += int(amount)
            else:
                t["awaiting_payment_count"] += n
        elif status == OrderStatus.REJECTED.value:
            t["rejected_orders"] += n
    passes_sold = await session.scalar(
        select(func.count()).select_from(Pass).join(Order, Order.id == Pass.order_id)
        .where(Pass.event_id == ev.id, Order.status == APPROVED, Pass.status == PassStatus.ACTIVE.value))
    passes_scanned = await session.scalar(
        select(func.count(func.distinct(PassScan.pass_id))).join(Pass, Pass.id == PassScan.pass_id).where(Pass.event_id == ev.id))

    # --- by tier
    tiers = (await session.scalars(select(PassTier).where(PassTier.event_id == ev.id).order_by(PassTier.sort_order, PassTier.price))).all()
    usage = await tier_usage(session, [x.id for x in tiers])
    collected = dict((await session.execute(
        select(OrderItem.tier_id, func.coalesce(func.sum(OrderItem.qty * OrderItem.unit_price), 0))
        .join(Order, Order.id == OrderItem.order_id).where(Order.event_id == ev.id, Order.status == APPROVED)
        .group_by(OrderItem.tier_id))).all())
    scanned_by_tier = dict((await session.execute(
        select(Pass.tier_id, func.count(func.distinct(Pass.id))).join(PassScan, PassScan.pass_id == Pass.id)
        .where(Pass.event_id == ev.id).group_by(Pass.tier_id))).all())
    by_tier = [TierRow(tier_id=x.id, name=x.name, price=x.price, capacity=x.capacity, sold=usage[x.id]["approved"],
                       pending=usage[x.id]["pending"], collected=int(collected.get(x.id, 0)),
                       scanned=scanned_by_tier.get(x.id, 0)) for x in tiers]

    # --- by wing (buyer's wing at report time)
    wing = func.coalesce(func.nullif(Membership.wing, ""), "—")
    wing_rows = (await session.execute(
        select(wing, func.count(func.distinct(Order.id)), func.coalesce(func.sum(OrderItem.qty), 0),
               func.coalesce(func.sum(OrderItem.qty * OrderItem.unit_price), 0))
        .select_from(Order).join(OrderItem, OrderItem.order_id == Order.id)
        .outerjoin(Membership, (Membership.user_id == Order.user_id) & (Membership.society_id == Order.society_id))
        .where(Order.event_id == ev.id, Order.status == APPROVED).group_by(wing).order_by(wing))).all()
    by_wing = [WingRow(wing=w, orders=o, passes=int(p), collected=int(c)) for w, o, p, c in wing_rows]

    days = await day_stats(session, ev.id, list(ev.days))
    return EventReport(
        event_id=ev.id, title=ev.title,
        totals=ReportTotals(**t, passes_sold=passes_sold, passes_scanned=passes_scanned),
        by_tier=by_tier, by_wing=by_wing, by_day=[DayRow(date=d["date"], scanned=d["scanned"], expected=d["expected"]) for d in days],
    )


def safe_cell(v):
    """Neutralise spreadsheet formula injection from user-controlled strings (names, flat numbers...)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


ORDER_COLUMNS = ["Order ref", "Status", "Placed at", "Member", "Phone", "Wing", "Flat", "Passes", "Amount (₹)", "UTR",
                 "Paid amount (₹)", "Flags", "Reviewed at", "Reject reason"]


async def order_rows(session: AsyncSession, ev: Event, limit: int = 50_000) -> list[list]:
    rows = (await session.execute(
        select(Order, User, Membership).join(User, User.id == Order.user_id)
        .outerjoin(Membership, (Membership.user_id == Order.user_id) & (Membership.society_id == Order.society_id))
        .where(Order.event_id == ev.id).options(selectinload(Order.items).selectinload(OrderItem.tier))
        .order_by(Order.created_at).limit(limit))).unique().all()
    out = []
    for o, u, m in rows:
        out.append([
            o.ref_code, o.status, o.created_at.astimezone(UTC).strftime("%Y-%m-%d %H:%M"), u.name or "", u.phone,
            (m.wing if m else "") or "", (m.flat_no if m else "") or "",
            ", ".join(f"{i.qty}× {i.tier.name}" for i in o.items), o.total_amount, o.utr or "",
            o.paid_amount if o.paid_amount is not None else "", ", ".join(o.flags),
            o.reviewed_at.astimezone(UTC).strftime("%Y-%m-%d %H:%M") if o.reviewed_at else "", o.reject_reason or "",
        ])
    return out


async def export_csv(session: AsyncSession, ev: Event) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(ORDER_COLUMNS)
    for r in await order_rows(session, ev):
        w.writerow([safe_cell(c) for c in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM → Excel opens ₹ / Gujarati / Hindi correctly


async def export_xlsx(session: AsyncSession, ev: Event, report: EventReport) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F46E5")

    def sheet(ws, header, rows):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append([safe_cell(c) for c in r])
        for i, col in enumerate(ws.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value or "")) for c in col) + 2, 48)
        ws.freeze_panes = "A2"

    t = report.totals
    ws = wb.active
    ws.title = "Summary"
    sheet(ws, ["Metric", "Value"], [
        ["Event", ev.title], ["Generated at", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")],
        ["Collected (₹)", t.collected_amount], ["Pending review (₹)", t.pending_amount],
        ["Pending review (orders)", t.pending_count], ["Awaiting payment (orders)", t.awaiting_payment_count],
        ["Approved orders", t.approved_orders], ["Rejected orders", t.rejected_orders],
        ["Passes sold", t.passes_sold], ["Passes scanned", t.passes_scanned],
    ])
    sheet(wb.create_sheet("By tier"), ["Tier", "Price (₹)", "Capacity", "Sold", "Pending", "Collected (₹)", "Scanned"],
          [[x.name, x.price, x.capacity, x.sold, x.pending, x.collected, x.scanned] for x in report.by_tier])
    sheet(wb.create_sheet("By wing"), ["Wing", "Orders", "Passes", "Collected (₹)"],
          [[x.wing, x.orders, x.passes, x.collected] for x in report.by_wing])
    sheet(wb.create_sheet("By day"), ["Date", "Scanned", "Expected"],
          [[x.date.isoformat(), x.scanned, x.expected] for x in report.by_day])
    sheet(wb.create_sheet("Orders"), ORDER_COLUMNS, await order_rows(session, ev))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
